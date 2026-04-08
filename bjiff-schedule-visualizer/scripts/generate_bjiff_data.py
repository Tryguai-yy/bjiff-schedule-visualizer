from __future__ import annotations

import html
import json
import re
import sys
import subprocess
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_XLSX = ROOT / "bjiff_schedule.xlsx"
OUTPUT_JS = ROOT / "data.js"
DOUBAN_CACHE = ROOT / "douban_cache.json"
MOBILE_USER_AGENT = (
    "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) "
    "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Mobile/15E148 Safari/604.1"
)
DESKTOP_USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36"
)
LOOKUP_VERSION = "v3"
TITLE_QUERY_ALIASES = {
    "当哈利遇上莎莉": ["当哈利遇到莎莉"],
    "当哈利遇上莎莉 4K": ["当哈利遇到莎莉"],
}
MANUAL_DOUBAN_FALLBACKS = {
    "机器人总动员": {
        "id": "2131459",
        "url": "https://m.douban.com/movie/subject/2131459/",
        "title": "机器人总动员",
        "originalTitle": "WALL·E",
        "year": "2008",
        "meta": "美国 / 科幻 / 动画 / 冒险 / 2008-06-27(美国)上映 / 片长98分钟",
        "poster": "https://qnmob3.doubanio.com/view/photo/large/public/p1461851991.jpg?imageView2/1/q/60/w/300/h/300/format/jpg",
        "rating": "9.3",
        "ratingCount": "1487731",
        "summary": "公元2805年，人类文明高度发展，却因污染和生活垃圾大量增加使得地球不再适于人类居住。地球人被迫乘坐飞船离开故乡，进行一次漫长无边的宇宙之旅。临行前他们委托Buynlarge的公司对地球垃圾进行清理，该公司开发了名为WALL·E的机器人担当此重任。",
        "matchScore": 999,
    },
    "当哈利遇上莎莉": {
        "id": "1291842",
        "url": "https://m.douban.com/movie/subject/1291842/",
        "title": "当哈利遇到莎莉",
        "originalTitle": "When Harry Met Sally...",
        "year": "1989",
        "meta": "美国 / 剧情 / 喜剧 / 爱情 / 1989-07-21(美国)上映 / 片长96分钟",
        "poster": "",
        "rating": "8.3",
        "ratingCount": "156671",
        "summary": "",
        "matchScore": 999,
    },
    "穿普拉达的女王": {
        "id": "1482072",
        "url": "https://m.douban.com/movie/subject/1482072/",
        "title": "穿普拉达的女王",
        "originalTitle": "The Devil Wears Prada",
        "year": "2006",
        "meta": "美国 / 法国 / 剧情 / 喜剧 / 2006-06-30(美国)上映 / 片长109分钟",
        "poster": "",
        "rating": "8.2",
        "ratingCount": "816858",
        "summary": "",
        "matchScore": 999,
    },
    "恐怖分子": {
        "id": "1305261",
        "url": "https://m.douban.com/movie/subject/1305261/",
        "title": "恐怖分子",
        "originalTitle": "恐怖份子",
        "year": "1986",
        "meta": "中国台湾 / 剧情 / 犯罪 / 1986-12-19(中国台湾)上映 / 片长109分钟",
        "poster": "",
        "rating": "9.0",
        "ratingCount": "94480",
        "summary": "",
        "matchScore": 999,
    },
    "史崔特先生的故事": {
        "id": "1298506",
        "url": "https://m.douban.com/movie/subject/1298506/",
        "title": "史崔特先生的故事",
        "originalTitle": "The Straight Story",
        "year": "1999",
        "meta": "法国 / 英国 / 美国 / 剧情 / 传记 / 1999-05-21(戛纳电影节)上映 / 片长112分钟",
        "poster": "",
        "rating": "8.5",
        "ratingCount": "41351",
        "summary": "",
        "matchScore": 999,
    },
}


def slugify(value: str) -> str:
    normalized = re.sub(r"[^\w\u4e00-\u9fff]+", "-", value.strip().lower())
    normalized = re.sub(r"-{2,}", "-", normalized).strip("-")
    return normalized or "item"


def compact_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u3000", " ")).strip()


def normalize_compare(value: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", (value or "").lower())


def clean_lookup_title(value: str) -> str:
    text = compact_spaces(str(value or ""))
    text = re.sub(r"\s*[（(]露天专场[)）]\s*$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*露天专场$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*(4k|imax)$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*(导演剪辑版|最终剪辑版|加长版)$", "", text)
    return compact_spaces(text)


def parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        text = compact_spaces(value)
        for fmt in ("%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
    return None


def load_douban_cache() -> dict[str, Any]:
    if DOUBAN_CACHE.exists():
        try:
            return json.loads(DOUBAN_CACHE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def build_cache_key(version: str, film: dict[str, Any]) -> str:
    return slugify(
        f"douban-{version}-{clean_lookup_title(film['title'])}-{clean_lookup_title(film.get('englishTitle', ''))}-{film.get('year', '')}"
    )


def build_legacy_cache_keys(film: dict[str, Any]) -> list[str]:
    keys = [build_cache_key("v2", film), build_cache_key("v1", film)]
    keys.append(slugify(f"douban-{film['title']}-{film.get('englishTitle', '')}-{film.get('year', '')}"))
    keys.append(slugify(f"douban-{clean_lookup_title(film['title'])}-{clean_lookup_title(film.get('englishTitle', ''))}-{film.get('year', '')}"))
    return keys


def find_legacy_cached_douban(cache: dict[str, Any], film: dict[str, Any]) -> dict[str, Any] | None:
    for key in build_legacy_cache_keys(film):
        cached = cache.get(key)
        if cached:
            return cached
    return None


def save_douban_cache(cache: dict[str, Any]) -> None:
    DOUBAN_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def curl_text(url: str, user_agent: str = MOBILE_USER_AGENT) -> str:
    try:
        result = subprocess.run(
            ["curl", "-L", "-A", user_agent, url],
            check=False,
            capture_output=True,
            text=True,
            timeout=12,
        )
    except subprocess.TimeoutExpired:
        return ""

    if result.returncode != 0:
        return ""
    return result.stdout


def extract_tag_content(html_text: str, tag: str, class_name: str) -> str:
    match = re.search(
        rf"<{tag}[^>]+class=\"{re.escape(class_name)}\"[^>]*>(.*?)</{tag}>",
        html_text,
        re.IGNORECASE | re.DOTALL,
    )
    if not match:
        return ""
    return compact_spaces(re.sub(r"<[^>]+>", " ", html.unescape(match.group(1))))


def parse_douban_search_results(html_text: str) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for match in re.finditer(
        r"<li>\s*<a href=\"/movie/subject/(?P<id>\d+)/\">(?P<body>.*?)</a>\s*</li>",
        html_text,
        re.IGNORECASE | re.DOTALL,
    ):
        body = match.group("body")
        title_match = re.search(r"<span class=\"subject-title\">(.*?)</span>", body, re.IGNORECASE | re.DOTALL)
        if not title_match:
            continue

        poster_match = re.search(r"<img src=\"([^\"]+)\"", body, re.IGNORECASE)
        rating_match = re.search(r"<p class=\"rating\">.*?<span>([\d.]+)</span>", body, re.IGNORECASE | re.DOTALL)
        results.append(
            {
                "id": match.group("id"),
                "title": compact_spaces(html.unescape(title_match.group(1))),
                "poster": poster_match.group(1) if poster_match else "",
                "rating": rating_match.group(1) if rating_match else "",
            }
        )

    return results


def fetch_bing_douban_candidates(query: str) -> list[dict[str, Any]]:
    search_url = f"https://www.bing.com/search?q={quote(f'{query} site:movie.douban.com/subject')}"
    html_text = curl_text(search_url, DESKTOP_USER_AGENT)
    results: list[dict[str, Any]] = []
    seen_ids: set[str] = set()

    for match in re.finditer(
        r"<h2[^>]*>\s*<a[^>]+href=\"https://movie\.douban\.com/subject/(?P<id>\d+)/\"[^>]*>(?P<title>.*?)</a>",
        html_text,
        re.IGNORECASE | re.DOTALL,
    ):
        subject_id = match.group("id")
        if subject_id in seen_ids:
            continue
        seen_ids.add(subject_id)
        title = compact_spaces(re.sub(r"<[^>]+>", " ", html.unescape(match.group("title"))))
        title = title.replace("(豆瓣)", "").strip()
        results.append(
            {
                "id": subject_id,
                "title": title,
                "poster": "",
                "rating": "",
            }
        )

    return results


def fetch_douban_search_results(query: str) -> list[dict[str, Any]]:
    if not query:
        return []
    url = f"https://m.douban.com/search/?query={quote(query)}"
    html_text = curl_text(url)
    results = parse_douban_search_results(html_text)
    if results:
        return results
    return fetch_bing_douban_candidates(query)


def extract_meta_content(html_text: str, attr_name: str, attr_value: str) -> str:
    patterns = [
        rf'<meta[^>]+{attr_name}="{re.escape(attr_value)}"[^>]+content="([^"]+)"',
        rf'<meta[^>]+content="([^"]+)"[^>]+{attr_name}="{re.escape(attr_value)}"',
    ]
    for pattern in patterns:
        match = re.search(pattern, html_text, re.IGNORECASE)
        if match:
            return html.unescape(match.group(1)).strip()
    return ""


def fetch_douban_detail(subject_id: str) -> dict[str, Any] | None:
    html_text = curl_text(f"https://m.douban.com/movie/subject/{subject_id}/")
    if not html_text:
        return None

    title = (
        extract_tag_content(html_text, "div", "sub-title")
        or extract_meta_content(html_text, "property", "og:title")
        or extract_meta_content(html_text, "itemprop", "name")
    )
    original_title = extract_tag_content(html_text, "div", "sub-original-title")
    meta = extract_tag_content(html_text, "div", "sub-meta")
    description = extract_meta_content(html_text, "name", "description")
    rating = extract_meta_content(html_text, "itemprop", "ratingValue")
    rating_count = extract_meta_content(html_text, "itemprop", "reviewCount")
    poster = extract_meta_content(html_text, "itemprop", "image") or extract_meta_content(html_text, "property", "og:image")

    summary = description
    if "简介：" in summary:
        summary = summary.split("简介：", 1)[1]
    summary = compact_spaces(summary)
    summary = re.sub(r"^[^。]*豆瓣评分[:：]\s*[\d.]+\s*", "", summary)

    year = ""
    year_match = re.search(r"[（(](\d{4})[)）]\s*$", original_title)
    if year_match:
        year = year_match.group(1)
        original_title = compact_spaces(re.sub(r"[（(]\d{4}[)）]\s*$", "", original_title))
    elif meta:
        meta_year_match = re.search(r"(\d{4})(?:-\d{2}-\d{2})?", meta)
        if meta_year_match:
            year = meta_year_match.group(1)

    title = title.replace(" - 电影", "").strip() if title else ""

    if not any([title, original_title, rating, poster, summary]):
        return None

    return {
        "id": subject_id,
        "url": f"https://m.douban.com/movie/subject/{subject_id}/",
        "title": title,
        "originalTitle": original_title,
        "year": year,
        "meta": meta,
        "poster": poster,
        "rating": rating,
        "ratingCount": rating_count,
        "summary": summary,
    }


def score_douban_candidate(film: dict[str, Any], candidate: dict[str, Any]) -> int:
    score = 0
    film_cn = normalize_compare(clean_lookup_title(film["title"]))
    film_en = normalize_compare(film.get("englishTitle", ""))
    candidate_cn = normalize_compare(str(candidate.get("title", "")))
    candidate_en = normalize_compare(str(candidate.get("originalTitle", "")))
    candidate_year = str(candidate.get("year", "")).strip()
    film_year = str(film.get("year", "")).strip()

    if film_cn and candidate_cn == film_cn:
        score += 120
    elif film_cn and (film_cn in candidate_cn or candidate_cn in film_cn):
        score += 70

    if film_en and candidate_en == film_en:
        score += 90
    elif film_en and candidate_en and (film_en in candidate_en or candidate_en in film_en):
        score += 45

    if film_year and candidate_year == film_year:
        score += 30
    elif film_year and candidate_year:
        try:
            if abs(int(candidate_year) - int(film_year)) <= 1:
                score += 15
        except ValueError:
            pass

    if candidate.get("rating"):
        score += 4

    return score


def lookup_douban(film: dict[str, Any]) -> dict[str, Any] | None:
    cleaned_title = clean_lookup_title(film["title"])
    cleaned_english = clean_lookup_title(film.get("englishTitle", ""))
    film_year = str(film.get("year", "")).strip()
    queries = []
    alias_queries = TITLE_QUERY_ALIASES.get(film["title"], []) + TITLE_QUERY_ALIASES.get(cleaned_title, [])
    for value in [
        cleaned_title,
        f"{cleaned_title} {film_year}".strip(),
        cleaned_english,
        f"{cleaned_english} {film_year}".strip(),
        *alias_queries,
    ]:
        value = compact_spaces(str(value or ""))
        if value and value not in queries:
            queries.append(value)

    candidates: dict[str, dict[str, Any]] = {}
    for query in queries[:3]:
        for candidate in fetch_douban_search_results(query):
            candidate_id = str(candidate.get("id", ""))
            if not candidate_id:
                continue
            current = candidates.get(candidate_id, candidate)
            current["searchScore"] = max(current.get("searchScore", 0), score_douban_candidate(film, candidate))
            candidates[candidate_id] = current
        if any(candidate.get("searchScore", 0) >= 120 for candidate in candidates.values()):
            break

    if not candidates:
        return None

    ranked_candidates = sorted(candidates.values(), key=lambda item: item.get("searchScore", 0), reverse=True)
    candidate_limit = 1 if ranked_candidates and ranked_candidates[0].get("searchScore", 0) >= 120 else 2
    best_detail = None
    best_score = 0

    for candidate in ranked_candidates[:candidate_limit]:
        detail = fetch_douban_detail(str(candidate["id"]))
        if not detail:
            continue

        detail["poster"] = detail["poster"] or candidate.get("poster", "")
        detail["rating"] = detail["rating"] or candidate.get("rating", "")
        detail["matchScore"] = score_douban_candidate(film, detail)
        if detail["matchScore"] > best_score:
            best_detail = detail
            best_score = detail["matchScore"]
        if best_score >= 150:
            break

    if not best_detail or best_score < 70:
        return MANUAL_DOUBAN_FALLBACKS.get(cleaned_title)

    return best_detail


def enrich_films_with_douban(films: list[dict[str, Any]]) -> list[dict[str, Any]]:
    cache = load_douban_cache()
    changed = False

    for film in films:
        cache_key = build_cache_key(LOOKUP_VERSION, film)
        cached = cache.get(cache_key)
        if cached:
            film["douban"] = cached
            continue

        legacy_cached = find_legacy_cached_douban(cache, film)
        if legacy_cached:
            cache[cache_key] = legacy_cached
            film["douban"] = legacy_cached
            changed = True
            continue

        douban = lookup_douban(film)
        cache[cache_key] = douban
        film["douban"] = douban
        changed = True
        time.sleep(0.12)

    if changed:
        save_douban_cache(cache)

    return films


def parse_main_sheet(workbook) -> list[dict[str, Any]]:
    ws = workbook["北京展映"]
    films: dict[tuple[str, str], dict[str, Any]] = {}

    for row in ws.iter_rows(min_row=4, values_only=True):
        unit, chinese_title, english_title, year, runtime, price, screening_at, cinema, hall, activity = row[:10]
        screening_dt = parse_datetime(screening_at)
        if not chinese_title or not screening_dt:
            continue

        chinese_title = compact_spaces(str(chinese_title))
        english_title = compact_spaces(str(english_title or ""))
        key = (chinese_title, english_title)

        if key not in films:
            films[key] = {
                "id": slugify(f"{chinese_title}-{english_title or year or 'film'}"),
                "title": chinese_title,
                "englishTitle": english_title,
                "unit": compact_spaces(str(unit or "未分类")),
                "year": int(year) if isinstance(year, (int, float)) else year,
                "runtime": int(runtime) if isinstance(runtime, (int, float)) else runtime,
                "price": int(price) if isinstance(price, (int, float)) else price,
                "isImmersive": False,
                "languages": "",
                "notes": [],
                "screenings": [],
            }

        film = films[key]
        if activity and activity not in film["notes"]:
            film["notes"].append(compact_spaces(str(activity)))

        film["screenings"].append(
            {
                "start": screening_dt.isoformat(),
                "dateLabel": screening_dt.strftime("%m.%d"),
                "weekdayLabel": screening_dt.strftime("%a"),
                "timeLabel": screening_dt.strftime("%H:%M"),
                "cinema": compact_spaces(str(cinema or "")),
                "hall": compact_spaces(str(hall or "")),
                "activity": compact_spaces(str(activity or "")),
            }
        )

    return finalize_films(list(films.values()))


def parse_immersive_sheet(workbook) -> list[dict[str, Any]]:
    ws = workbook["无界∞沉浸单元"]
    column_groups: dict[int, str] = {}
    current_group = "无界∞沉浸单元"

    for idx, cell in enumerate(ws[2], start=1):
        if idx == 1:
            continue
        if cell.value:
            current_group = compact_spaces(str(cell.value))
        column_groups[idx] = current_group

    films: list[dict[str, Any]] = []
    film_columns: list[tuple[int, dict[str, Any]]] = []

    for idx in range(2, ws.max_column + 1):
        raw = ws.cell(row=3, column=idx).value
        if not raw:
            continue
        film = parse_immersive_header(str(raw), column_groups.get(idx, "无界∞沉浸单元"), idx)
        films.append(film)
        film_columns.append((idx, film))

    for row_idx in range(4, ws.max_row + 1):
        raw_date = ws.cell(row=row_idx, column=1).value
        date_label = normalize_immersive_date(raw_date)
        if not date_label:
            continue

        for col_idx, film in film_columns:
            raw_times = ws.cell(row=row_idx, column=col_idx).value
            times = parse_immersive_times(str(raw_times or ""))
            for time_label in times:
                film["screenings"].append(
                    {
                        "start": f"2026-{date_label.replace('.', '-') }T{time_label}:00",
                        "dateLabel": date_label,
                        "weekdayLabel": "",
                        "timeLabel": time_label,
                        "cinema": "中国传媒大学",
                        "hall": "无界∞沉浸单元",
                        "activity": "",
                    }
                )

    return finalize_films(films)


def parse_immersive_header(raw: str, group_name: str, index: int) -> dict[str, Any]:
    cleaned = raw.replace("\r", "\n").replace("（", "(").replace("）", ")")
    lines = [compact_spaces(line) for line in cleaned.split("\n") if compact_spaces(line)]
    joined = " ".join(lines)

    price_match = re.search(r"\((\d+)元/场\)|（(\d+)元/场）", cleaned)
    runtime_match = re.search(r"(\d+)(?:分钟|分钟/场|分钟10秒/场|分钟/场)", cleaned)

    language_line = ""
    for line in lines[1:]:
        if any(token in line for token in ["对白", "字幕", "语", "手语"]):
            language_line = line
            break

    title_line = lines[0] if lines else f"沉浸单元 #{index}"
    title_line = title_line.replace("、 水·体", "、水·体")

    split_match = re.match(r"(?P<cn>[\u4e00-\u9fff·：:《》“”‘’、·A-Za-z0-9\s\-/'&]+?)(?P<en>[A-Z][A-Za-z0-9\s\-/'&:,!.·]+)$", title_line)
    if split_match:
        chinese_title = compact_spaces(split_match.group("cn"))
        english_title = compact_spaces(split_match.group("en"))
    else:
        fallback = re.split(r"(?=[A-Z][a-z])", title_line, maxsplit=1)
        if len(fallback) == 2:
            chinese_title, english_title = map(compact_spaces, fallback)
        else:
            chinese_title, english_title = compact_spaces(title_line), ""

    return {
        "id": slugify(f"immersive-{index}-{chinese_title}"),
        "title": chinese_title,
        "englishTitle": english_title,
        "unit": f"无界∞沉浸单元 · {group_name}",
        "year": 2026,
        "runtime": int(runtime_match.group(1)) if runtime_match else None,
        "price": int(next(group for group in price_match.groups() if group)) if price_match else None,
        "isImmersive": True,
        "languages": language_line,
        "notes": [joined],
        "screenings": [],
    }


def normalize_immersive_date(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, (int, float)):
        month = int(raw)
        day = int(round((float(raw) - month) * 100))
        if day:
            return f"{month:02d}.{day:02d}"
    value = compact_spaces(str(raw))
    match = re.search(r"(\d{1,2})\.(\d{1,2})", value)
    if match:
        return f"{int(match.group(1)):02d}.{int(match.group(2)):02d}"
    return ""


def parse_immersive_times(raw: str) -> list[str]:
    raw = raw.replace("：", ":").replace(";", ":")
    return re.findall(r"\d{1,2}:\d{2}", raw)


def finalize_films(films: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for film in films:
        film["screenings"].sort(key=lambda item: item["start"])
        unique_venues = []
        seen = set()
        schedule_by_date: defaultdict[str, list[str]] = defaultdict(list)
        creator_event_count = 0
        for screening in film["screenings"]:
            venue = screening["cinema"]
            if venue and venue not in seen:
                seen.add(venue)
                unique_venues.append(venue)
            schedule_by_date[screening["dateLabel"]].append(screening["timeLabel"])
            if any(token in screening["activity"] for token in ["映后", "主创"]):
                creator_event_count += 1

        daily_schedule = []
        for date_label in sorted(schedule_by_date):
            unique_times = sorted(set(schedule_by_date[date_label]))
            daily_schedule.append(
                {
                    "dateLabel": date_label,
                    "times": unique_times,
                    "count": len(unique_times),
                }
            )

        film["venues"] = unique_venues
        film["screeningCount"] = len(film["screenings"])
        film["sessionCount"] = len(film["screenings"])
        film["dayCount"] = len(daily_schedule)
        film["scheduleByDate"] = daily_schedule
        film["firstScreening"] = film["screenings"][0]["start"] if film["screenings"] else None
        film["lastScreening"] = film["screenings"][-1]["start"] if film["screenings"] else None
        film["creatorEventCount"] = creator_event_count
        film["hasCreatorEvent"] = creator_event_count > 0 or any(
            token in note for note in film["notes"] for token in ["映后", "主创"]
        )

        if film["isImmersive"]:
            average_daily_slots = round(sum(day["count"] for day in daily_schedule) / len(daily_schedule)) if daily_schedule else 0
            film["averageDailySlots"] = average_daily_slots
            film["displayScreeningCount"] = film["dayCount"]
            film["displayScreeningLabel"] = f"{film['dayCount']} 天循环"
            film["scheduleSummary"] = (
                f"{film['dayCount']} 天循环放映，每天约 {average_daily_slots} 个体验时段"
                if average_daily_slots
                else "循环放映"
            )
        else:
            film["averageDailySlots"] = 0
            film["displayScreeningCount"] = film["screeningCount"]
            film["displayScreeningLabel"] = f"{film['screeningCount']} 场"
            film["scheduleSummary"] = f"{film['screeningCount']} 场放映，横跨 {film['dayCount']} 天"
    return films


def build_payload(main_films: list[dict[str, Any]], immersive_films: list[dict[str, Any]]) -> dict[str, Any]:
    films = main_films + immersive_films
    units = sorted({film["unit"] for film in films})
    dates = sorted({screening["dateLabel"] for film in films for screening in film["screenings"]})
    venues = sorted({screening["cinema"] for film in films for screening in film["screenings"] if screening["cinema"]})

    stats = {
        "filmCount": len(films),
        "screeningCount": sum(film["sessionCount"] for film in main_films),
        "allSessionCount": sum(film["sessionCount"] for film in films),
        "immersiveFilmCount": len(immersive_films),
        "immersiveSessionCount": sum(film["sessionCount"] for film in immersive_films),
        "creatorEventFilmCount": sum(1 for film in films if film["hasCreatorEvent"]),
        "doubanMatchedFilmCount": sum(1 for film in films if film.get("douban")),
        "immersiveDayCount": len(
            {
                day["dateLabel"]
                for film in immersive_films
                for day in film["scheduleByDate"]
            }
        ),
        "unitCount": len(units),
        "venueCount": len(venues),
    }

    return {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "festival": {
            "name": "第十六届北京国际电影节",
            "subtitle": "北京展映排片可视化",
            "dateRange": "2026.04.16 - 2026.04.26",
        },
        "stats": stats,
        "units": units,
        "dates": dates,
        "venues": venues,
        "films": sorted(films, key=lambda item: (item["firstScreening"] or "", item["title"])),
    }


def resolve_source_xlsx() -> Path:
    if len(sys.argv) > 1:
        source = Path(sys.argv[1]).expanduser()
    else:
        source = DEFAULT_SOURCE_XLSX

    if not source.is_absolute():
        source = (ROOT / source).resolve()

    return source


def main() -> None:
    source_xlsx = resolve_source_xlsx()
    if not source_xlsx.exists():
        raise FileNotFoundError(
            f"Excel file not found: {source_xlsx}\n"
            f"Place it at {DEFAULT_SOURCE_XLSX.name} in the project root, or pass the path explicitly."
        )

    workbook = load_workbook(source_xlsx, data_only=True)
    main_films = parse_main_sheet(workbook)
    immersive_films = parse_immersive_sheet(workbook)
    payload = build_payload(
        enrich_films_with_douban(main_films),
        enrich_films_with_douban(immersive_films),
    )
    OUTPUT_JS.write_text(
        "window.BJIFF_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT_JS}")


if __name__ == "__main__":
    main()
